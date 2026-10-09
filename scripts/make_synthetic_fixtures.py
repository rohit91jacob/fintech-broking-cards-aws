"""Generate SYNTHETIC NSE and RBI fixtures in the real file layouts.

NSE and RBI do not permit redistributing their files, so the repository cannot ship real
samples. These fixtures copy the *layout* (column names, header rows, units, quirks such as
the stray cell in front of a bank row) with made-up numbers and names, so the parsers and
the whole pipeline can be tested offline. Every symbol is ``SYN..``, every ISIN ``INSYN...``,
every bank ``SYNTHETIC ... BANK`` and every workbook says SYNTHETIC in its title.

    uv run python scripts/make_synthetic_fixtures.py --out tests/fixtures/raw
"""

from __future__ import annotations

import argparse
import csv
import io
import random
import zipfile
from datetime import date, timedelta
from pathlib import Path

import openpyxl

SESSIONS = [date(2026, 9, 28) + timedelta(days=i) for i in range(5)]  # Mon 28 Sep .. Fri 2 Oct
HOLIDAY = date(2026, 10, 2)  # Gandhi Jayanti: no session, like the real calendar
NSE_COLUMNS = [
    "TradDt",
    "BizDt",
    "Sgmt",
    "Src",
    "FinInstrmTp",
    "FinInstrmId",
    "ISIN",
    "TckrSymb",
    "SctySrs",
    "XpryDt",
    "FininstrmActlXpryDt",
    "StrkPric",
    "OptnTp",
    "FinInstrmNm",
    "OpnPric",
    "HghPric",
    "LwPric",
    "ClsPric",
    "LastPric",
    "PrvsClsgPric",
    "UndrlygPric",
    "SttlmPric",
    "OpnIntrst",
    "ChngInOpnIntrst",
    "TtlTradgVol",
    "TtlTrfVal",
    "TtlNbOfTxsExctd",
    "SsnId",
    "NewBrdLotQty",
    "Rmks",
    "Rsvd1",
    "Rsvd2",
    "Rsvd3",
    "Rsvd4",
]
MONTHS = [
    "January",
    "February",
    "March",
    "April",
    "May",
    "June",
    "July",
    "August",
    "September",
    "October",
    "November",
    "December",
]


def nse(out: Path, rng: random.Random) -> None:
    d = out / "nse"
    d.mkdir(parents=True, exist_ok=True)
    symbols = [(f"SYN{i:02d}", f"INSYN{i:05d}0", rng.uniform(80, 4000)) for i in range(1, 41)]
    prev = {s: p for s, _, p in symbols}
    for day in SESSIONS:
        if day == HOLIDAY:
            continue
        buf = io.StringIO()
        w = csv.DictWriter(buf, NSE_COLUMNS)
        w.writeheader()
        for i, (sym, isin, _) in enumerate(symbols):
            p0 = prev[sym]
            close = round(p0 * (1 + rng.gauss(0, 0.015)), 2)
            hi, lo = round(max(p0, close) * 1.01, 2), round(min(p0, close) * 0.99, 2)
            vol = rng.randint(50_000, 5_000_000) * (41 - i)
            vwap = (hi + lo + close) / 3
            w.writerow(
                {
                    "TradDt": day.isoformat(),
                    "BizDt": day.isoformat(),
                    "Sgmt": "CM",
                    "Src": "NSE",
                    "FinInstrmTp": "STK",
                    "FinInstrmId": 90000 + i,
                    "ISIN": isin,
                    "TckrSymb": sym,
                    "SctySrs": "EQ",
                    "FinInstrmNm": f"SYNTHETIC COMPANY {i:02d} LTD",
                    "OpnPric": p0,
                    "HghPric": hi,
                    "LwPric": lo,
                    "ClsPric": close,
                    "LastPric": close,
                    "PrvsClsgPric": p0,
                    "SttlmPric": close,
                    "TtlTradgVol": vol,
                    "TtlTrfVal": round(vol * vwap, 2),
                    "TtlNbOfTxsExctd": vol // 120,
                    "SsnId": "F1",
                    "NewBrdLotQty": 1,
                }
            )
            prev[sym] = close
        # one non-EQ row, which the pipeline must ignore
        w.writerow(
            {
                "TradDt": day.isoformat(),
                "BizDt": day.isoformat(),
                "Sgmt": "CM",
                "Src": "NSE",
                "FinInstrmTp": "STK",
                "ISIN": "INSYNGB00010",
                "TckrSymb": "SYNGB28",
                "SctySrs": "GB",
                "FinInstrmNm": "SYNTHETIC GOLD BOND",
                "OpnPric": 100,
                "HghPric": 101,
                "LwPric": 99,
                "ClsPric": 100,
                "PrvsClsgPric": 100,
                "TtlTradgVol": 10,
                "TtlTrfVal": 1000,
                "TtlNbOfTxsExctd": 2,
                "SsnId": "F1",
                "NewBrdLotQty": 1,
            }
        )
        name = f"BhavCopy_NSE_CM_0_0_0_{day:%Y%m%d}_F_0000"
        with zipfile.ZipFile(d / f"{name}.csv.zip", "w", zipfile.ZIP_DEFLATED) as zf:
            zf.writestr(f"{name}.csv", buf.getvalue())


def _month_back(m: date, n: int) -> date:
    y, mo = divmod(m.year * 12 + m.month - 1 - n, 12)
    return date(y, mo + 1, 1)


def psi(out: Path, release: date, rng: random.Random) -> None:
    d = out / "rbi" / "psi"
    d.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{MONTHS[release.month - 1]} {release.year}"
    ly, prev = _month_back(release, 12), _month_back(release, 1)

    def growth(m: date) -> float:  # synthetic trend: ~2 % a month
        return 1.02 ** ((m.year - 2025) * 12 + m.month)

    def header(row: int) -> None:
        ws.cell(row, 3, "FY 2025-26")
        ws.cell(row, 4, f"{ly.year}\n{MONTHS[ly.month - 1]}")
        ws.cell(row, 5, prev.year)
        ws.cell(row + 1, 5, MONTHS[prev.month - 1])
        ws.cell(row, 6, release.year)
        ws.cell(row + 1, 6, MONTHS[release.month - 1])

    ws.cell(1, 2, "Payment System Indicators (SYNTHETIC TEST FIXTURE - not RBI data)")
    ws.cell(2, 2, "PART I - Payment System Indicators - Payment & Settlement System Statistics")
    ws.cell(3, 3, "Volume (lakh)")
    ws.cell(3, 7, "Value (₹ crore)")
    header(4)
    rows = [
        ("2.3 IMPS", 4000, 650000),
        ("2.5 NEFT", 9000, 4500000),
        ("2.6 UPI @", 200000, 2600000),
        ("4 Card Payments", 6500, 240000),
        ("4.1 Credit Cards ", 5500, 200000),
        ("4.1.1 PoS based $", 2700, 77000),
        ("4.1.2 Others $", 2800, 123000),
        ("4.2 Debit Cards", 1000, 37000),
        ("Total Digital Payments (1+2+3+4+5)", 260000, 30000000),
    ]
    r = 7
    for label, vol, val in rows:
        ws.cell(r, 2, label)
        ws.cell(r, 3, vol * 12)
        ws.cell(r, 7, val * 12)
        for j, m in ((4, ly), (5, prev), (6, release)):
            ws.cell(r, j, round(vol * growth(m) * rng.uniform(0.99, 1.01), 2))
            ws.cell(r, j + 4, round(val * growth(m) * rng.uniform(0.99, 1.01), 2))
        r += 1
    r += 2
    ws.cell(r, 2, "PART III - Payment Infrastructures (lakh)")
    header(r + 1)
    for k, (label, n) in enumerate(
        (
            ("1 Number of Cards ", 11500),
            ("1.1 Credit Cards", 1200),
            ("1.2 Debit Cards", 10300),
            ("7 UPI QR *", 8000),
        )
    ):
        ws.cell(r + 4 + k, 2, label)
        for j, m in ((4, ly), (5, prev), (6, release)):
            ws.cell(r + 4 + k, j, round(n * (1 + 0.004 * ((m.year - 2025) * 12 + m.month)), 2))
    r += 10
    ws.cell(r, 2, "PART V - Domestic Payment Frauds")
    for k in range(3):
        m = _month_back(release, 2 - k)
        ws.cell(r + 2 + k, 2, f"{MONTHS[m.month - 1]} {m.year}")
        ws.cell(r + 2 + k, 3, 3.0 + k / 10)
        ws.cell(r + 2 + k, 4, 450.0 + k)
    wb.save(d / f"SYNTHETIC_PSI_{release:%Y_%m}.XLSX")


def cards(out: Path, month: date, rng: random.Random) -> None:
    d = out / "rbi" / "cards"
    d.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"{MONTHS[month.month - 1]} {month.year}"
    ws.cell(
        2,
        2,
        f"SYNTHETIC TEST FIXTURE - ATM, Acceptance Infrastructure and Card Statistics for the Month of {MONTHS[month.month - 1]} {month.year}",
    )
    ws.cell(3, 2, "Sr. No.")
    ws.cell(3, 3, "Bank Name")
    for col in range(4, 30):
        ws.cell(8, col, col - 3)  # the numbered header row 1..26
    groups = {
        "Public Sector Banks": 3,
        "Private Sector Banks": 4,
        "Foreign Banks": 2,
        "Payment Banks": 1,
        "Small Finance Banks": 2,
    }
    r, sr = 9, 1
    totals = [0.0] * 30
    for group, n in groups.items():
        ws.cell(r, 2, group)
        r += 1
        for _ in range(n):
            ccards = 0 if group == "Payment Banks" else rng.randint(200_000, 25_000_000)
            vals = {
                4: rng.randint(100, 9000),
                5: rng.randint(10, 9000),
                6: rng.randint(0, 2_000_000),
                10: ccards,
                11: rng.randint(1_000_000, 200_000_000),
            }
            if ccards:
                pos_v, ecom_v = (
                    int(ccards * rng.uniform(1.0, 3.0)),
                    int(ccards * rng.uniform(1.0, 3.5)),
                )
                vals.update(
                    {
                        12: pos_v,
                        13: round(pos_v * rng.uniform(2.0, 3.5), 3),
                        14: ecom_v,
                        15: round(ecom_v * rng.uniform(3.0, 6.0), 3),
                        16: 0,
                        17: 0,
                    }
                )
            if sr == 7:
                ws.cell(r, 1, " ")  # the stray cell seen in real releases
            ws.cell(r, 2, sr)
            ws.cell(r, 3, f"SYNTHETIC {group.split()[0].upper()} BANK {sr:02d}")
            for col, v in vals.items():
                ws.cell(r, col, v)
                totals[col] += v
            r += 1
            sr += 1
    ws.cell(r, 2, "Total")
    for col in range(4, 18):
        ws.cell(r, col, totals[col])
    ws.cell(r + 3, 1, "Note")
    ws.cell(r + 3, 2, "The data is SYNTHETIC")
    wb.save(d / f"SYNTHETIC_CARDS_{month:%Y_%m}.XLSX")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()
    rng = random.Random(20261009)
    nse(args.out, rng)
    for m in (date(2026, 6, 1), date(2026, 7, 1), date(2026, 8, 1)):
        psi(args.out, m, rng)
        cards(args.out, m, rng)
    for p in sorted(args.out.rglob("*")):
        if p.is_file() and ("nse" in p.parts or "rbi" in p.parts):
            print(f"{p.stat().st_size:>8} {p.relative_to(args.out)}")


if __name__ == "__main__":
    main()
