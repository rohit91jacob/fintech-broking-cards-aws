"""Reserve Bank of India payment statistics (monthly XLSX releases).

* **Payment System Indicators (PSI)**: system-wide volumes and values, including UPI, credit
  and debit card payments, the number of cards outstanding and reported payment frauds.
  Listing page: www.rbi.org.in/Scripts/PSIUserView.aspx.
* **ATM, Acceptance Infrastructure and Card Statistics**: bank-wise credit cards outstanding
  and credit/debit card spends (PoS, e-commerce, others). Listing: /Scripts/ATMView.aspx.

UPI figures here are the RBI's own series (RBI sources them from NPCI). NPCI's website blocks
automated clients, so it is not used directly.

The RBI's website terms prohibit caching or framing its content without permission, so the
files are downloaded at run time into the git-ignored raw directory and never committed. The
results site publishes only derived aggregates, with attribution. Tests use synthetic
workbooks in the same layout (tests/fixtures/synthetic/rbi).
"""

from __future__ import annotations

import calendar
import logging
import re
from dataclasses import dataclass
from datetime import date
from pathlib import Path
from typing import Any

import openpyxl
import pyarrow as pa

from fbc import http

log = logging.getLogger(__name__)

BASE = "https://www.rbi.org.in/Scripts"
LISTINGS = {
    "psi": (
        f"{BASE}/PSIUserView.aspx",
        re.compile(r"https://rbidocs\.rbi\.org\.in/rdocs/PSI/DOCs/[A-Z0-9]+\.XLSX", re.I),
    ),
    "cards": (
        f"{BASE}/ATMView.aspx",
        re.compile(r"https://rbidocs\.rbi\.org\.in/rdocs/ATM/DOCs/[A-Z0-9]+\.XLSX", re.I),
    ),
}
MONTHS = {name.lower(): i for i, name in enumerate(calendar.month_name) if name}
# RBI typos seen in real releases ("Jauuary 2025").
MONTHS["jauuary"] = 1


class LayoutError(Exception):
    """The workbook no longer has the layout this parser was written for."""


def fetch(raw_dir: Path, months: int, refresh: bool = True) -> dict[str, list[Path]]:
    sess = http.session()
    out: dict[str, list[Path]] = {}
    for kind, (listing, pattern) in LISTINGS.items():
        page = http.fetch(sess, listing, raw_dir / "rbi" / f"{kind}_listing.html", refresh=refresh)
        links = list(dict.fromkeys(pattern.findall(page.decode("utf-8", errors="replace"))))
        if not links:
            raise RuntimeError(f"no XLSX links found on {listing}")
        paths = []
        for url in links[:months]:
            path = raw_dir / "rbi" / kind / url.rsplit("/", 1)[1].upper()
            http.fetch(sess, url, path, expect=b"PK")  # releases are immutable once published
            paths.append(path)
        out[kind] = paths
    return out


def _month_from_text(text: str) -> date | None:
    m = re.search(r"([A-Za-z]+)\s*,?\s*(\d{4})", text)
    if not m or m.group(1).lower() not in MONTHS:
        return None
    return date(int(m.group(2)), MONTHS[m.group(1).lower()], 1)


def _cell_month(top: Any, bottom: Any) -> date | None:
    """Header cells are either ``'2025\\nAugust'`` or year in one row and month in the next."""
    if isinstance(top, str) and "\n" in top:
        year, month = top.split("\n", 1)
        return _month_from_text(f"{month.strip()} {year.strip()}")
    if isinstance(top, int | float) and isinstance(bottom, str):
        return _month_from_text(f"{bottom.strip()} {int(top)}")
    return None


# ------------------------------------------------------------------------------------- PSI

PSI_METRICS = {
    # label prefix -> metric name (Parts I and II: volume in lakh, value in crore)
    "2.6 UPI": "upi",
    "2.3 IMPS": "imps",
    "2.5 NEFT": "neft",
    "4 Card Payments": "card_payments",
    "4.1 Credit Cards": "credit_card_payments",
    "4.1.1 PoS based": "credit_card_pos",
    "4.1.2 Others": "credit_card_others",
    "4.2 Debit Cards": "debit_card_payments",
    "Total Digital Payments": "total_digital_payments",
}
PSI_COUNTS = {
    # Part III, number outstanding in lakh
    "1.1 Credit Cards": "credit_cards_outstanding",
    "1.2 Debit Cards": "debit_cards_outstanding",
    "7 UPI QR": "upi_qr_codes",
}

PSI_SCHEMA = pa.schema(
    [
        ("release_month", pa.date32()),
        ("period_month", pa.date32()),
        ("metric", pa.string()),
        ("measure", pa.string()),
        ("value", pa.float64()),
        ("source_file", pa.string()),
    ]
)


@dataclass
class _Part:
    name: str
    period_cols: dict[int, date]


def parse_psi(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    rows = [r + [None] * (16 - len(r)) for r in rows]
    release = _month_from_text(str(ws.title)) or _month_from_text(str(rows[0][1] or ""))
    out: list[dict[str, Any]] = []
    part: _Part | None = None
    for i, row in enumerate(rows):
        label = row[1].strip() if isinstance(row[1], str) else None
        if label and label.upper().startswith("PART "):
            part = _Part(label.split(" ", 2)[1].strip(" -"), {})
            continue
        if part is None:
            continue
        if not part.period_cols:
            # header row: column 3 is the same month last year, 4 and 5 the latest two months
            for j in (3, 4, 5):
                month = _cell_month(row[j], rows[i + 1][j] if i + 1 < len(rows) else None)
                if month:
                    part.period_cols[j] = month
            if part.name in ("V",) and label:
                part.period_cols = {}
            continue
        if not label:
            continue
        if part.name in ("I", "II"):
            metric = next((m for p, m in PSI_METRICS.items() if label.startswith(p)), None)
            if metric is None or (
                metric == "credit_card_payments" and not label.startswith("4.1 ")
            ):
                continue
            if metric == "card_payments" and not label.startswith("4 Card"):
                continue
            for j, month in part.period_cols.items():
                for measure, col in (("volume_lakh", j), ("value_crore", j + 4)):
                    if isinstance(row[col], int | float):
                        out.append(_psi_row(release, month, metric, measure, row[col], path))
        elif part.name == "III":
            metric = next((m for p, m in PSI_COUNTS.items() if label.startswith(p)), None)
            if metric is None:
                continue
            for j, month in part.period_cols.items():
                if isinstance(row[j], int | float):
                    out.append(_psi_row(release, month, metric, "count_lakh", row[j], path))
    out.extend(_parse_psi_frauds(rows, release, path))
    if not any(r["metric"] == "upi" for r in out):
        raise LayoutError(f"{path.name}: UPI row not found")
    return out


def _parse_psi_frauds(
    rows: list[list[Any]], release: date | None, path: Path
) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    in_part = False
    for row in rows:
        label = row[1].strip() if isinstance(row[1], str) else ""
        if label.upper().startswith("PART V"):
            in_part = True
            continue
        if not in_part:
            continue
        month = _month_from_text(label) if re.fullmatch(r"[A-Za-z]+ \d{4}", label) else None
        if month and isinstance(row[2], int | float):
            out.append(
                _psi_row(release, month, "domestic_payment_frauds", "volume_lakh", row[2], path)
            )
            out.append(
                _psi_row(release, month, "domestic_payment_frauds", "value_crore", row[3], path)
            )
    return out


def _psi_row(
    release: date | None, month: date, metric: str, measure: str, value: Any, path: Path
) -> dict[str, Any]:
    return {
        "release_month": release,
        "period_month": month,
        "metric": metric,
        "measure": measure,
        "value": float(value),
        "source_file": path.name,
    }


# ----------------------------------------------------------------------------- card statistics

# Column offsets in the bank-wise sheet (0-based), checked against the numbered header row.
CARD_COLUMNS = {
    "atms_onsite": 3,
    "atms_offsite": 4,
    "pos_terminals": 5,
    "micro_atms": 6,
    "bharat_qr": 7,
    "upi_qr": 8,
    "credit_cards_outstanding": 9,
    "debit_cards_outstanding": 10,
    "cc_pos_volume": 11,
    "cc_pos_value_thousand": 12,
    "cc_ecom_volume": 13,
    "cc_ecom_value_thousand": 14,
    "cc_others_volume": 15,
    "cc_others_value_thousand": 16,
    "cc_atm_cash_volume": 17,
    "cc_atm_cash_value_thousand": 18,
    "dc_pos_volume": 19,
    "dc_pos_value_thousand": 20,
    "dc_ecom_volume": 21,
    "dc_ecom_value_thousand": 22,
}

CARD_SCHEMA = pa.schema(
    [("period_month", pa.date32()), ("bank_group", pa.string()), ("bank_name", pa.string())]
    + [(name, pa.float64()) for name in CARD_COLUMNS]
    + [("is_total_row", pa.bool_()), ("source_file", pa.string())]
)

BANK_GROUPS = (
    "Public Sector Banks",
    "Private Sector Banks",
    "Foreign Banks",
    "Payment Banks",
    "Small Finance Banks",
)


def parse_card_stats(path: Path) -> list[dict[str, Any]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [list(r) + [None] * 32 for r in ws.iter_rows(values_only=True)]
    wb.close()
    title = next(
        (str(c) for r in rows[:6] for c in r if isinstance(c, str) and "Month of" in c), ""
    )
    month = _month_from_text(title.split("Month of", 1)[-1]) if title else None
    if month is None:
        raise LayoutError(f"{path.name}: month not found in title")
    numbered = next((r for r in rows[:15] if r[3] == 1 and r[4] == 2), None)
    if numbered is None or numbered[18] != 16:
        raise LayoutError(f"{path.name}: numbered header row not found or columns moved")

    out: list[dict[str, Any]] = []
    group = None
    for row in rows:
        label = row[1].strip() if isinstance(row[1], str) else None
        if label in BANK_GROUPS:
            group = label
            continue
        is_total = label == "Total"
        name = row[2].strip() if isinstance(row[2], str) else None
        if not is_total and not (isinstance(row[1], int | float) and name):
            continue
        rec: dict[str, Any] = {
            "period_month": month,
            "bank_group": None if is_total else group,
            "bank_name": "TOTAL" if is_total else name,
            "is_total_row": is_total,
            "source_file": path.name,
        }
        for col, idx in CARD_COLUMNS.items():
            value = row[idx]
            rec[col] = float(value) if isinstance(value, int | float) else None
        out.append(rec)
        if is_total:
            break
    if not any(r["is_total_row"] for r in out):
        raise LayoutError(f"{path.name}: Total row not found")
    return out


def parse_dir(rbi_dir: Path) -> tuple[pa.Table, pa.Table]:
    psi: list[dict[str, Any]] = []
    cards: list[dict[str, Any]] = []
    for path in sorted((rbi_dir / "psi").glob("*.XLSX")):
        psi.extend(parse_psi(path))
    for path in sorted((rbi_dir / "cards").glob("*.XLSX")):
        cards.extend(parse_card_stats(path))
    return pa.Table.from_pylist(psi, schema=PSI_SCHEMA), pa.Table.from_pylist(
        cards, schema=CARD_SCHEMA
    )
